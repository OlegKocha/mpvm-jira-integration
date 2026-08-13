"""Create formatted per-asset vulnerability reports in XLSX format."""

from __future__ import annotations

import os
import re
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ..models import Asset, localized_severity


XLSX_HEADERS = (
    "№",
    "CVE",
    "BDU",
    "ОС",
    "Пакет",
    "Версия пакета",
    "Описание",
    "How to fix",
    "CVSS vector",
    "Критичность",
    "CVSS score",
)
_ALLOWED_SEVERITIES = {
    "Критический",
    "Высокий",
    "Средний",
    "Низкий",
}
_COLUMN_WIDTHS = (7, 22, 20, 28, 28, 20, 60, 60, 42, 18, 13)
_ILLEGAL_XML_CHARACTERS = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")
_MAX_CELL_LENGTH = 32_767


def export_asset_xlsx(asset: Asset, directory: Path) -> Path:
    """Write a formatted asset vulnerability report and return its path."""
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f"vulnerabilities_{_safe_name(asset.title)}.xlsx"
    workbook = _build_workbook(asset)

    with tempfile.NamedTemporaryFile(
        dir=directory,
        prefix=".xlsx-",
        suffix=".xlsx",
        delete=False,
    ) as stream:
        temporary = Path(stream.name)

    try:
        workbook.save(temporary)
        os.replace(temporary, path)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise
    finally:
        workbook.close()
    return path


def _build_workbook(asset: Asset) -> Workbook:
    """Build an Excel workbook containing one vulnerability worksheet."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Уязвимости"
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False

    sheet.append(XLSX_HEADERS)
    for number, vulnerability in enumerate(asset.vulnerabilities, 1):
        severity = localized_severity(
            vulnerability.severity,
            vulnerability.cvss_score,
        )
        sheet.append(
            (
                number,
                _cell_text(", ".join(vulnerability.cves) or "-"),
                _cell_text(", ".join(vulnerability.bdus) or "-"),
                _cell_text(asset.operating_system or "-"),
                _cell_text(vulnerability.package_names or "-"),
                _cell_text(vulnerability.package_versions or "-"),
                _cell_text(vulnerability.description.strip() or "-"),
                _cell_text(vulnerability.remediation.strip() or "-"),
                _cell_text(vulnerability.cvss_vector.strip() or "-"),
                severity if severity in _ALLOWED_SEVERITIES else "-",
                vulnerability.cvss_score,
            )
        )

    _format_worksheet(sheet)
    return workbook


def _format_worksheet(sheet: Worksheet) -> None:
    """Apply readable Excel formatting to a vulnerability worksheet."""
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    border = Border(bottom=Side(style="thin", color="B4C6E7"))

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 24

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
        row[0].alignment = Alignment(horizontal="center", vertical="top")
        row[10].alignment = Alignment(horizontal="center", vertical="top")
        row[10].number_format = "0.0"

    for column, width in enumerate(_COLUMN_WIDTHS, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width

    sheet.auto_filter.ref = sheet.dimensions


def _cell_text(value: str) -> str:
    """Remove characters unsupported by XLSX and enforce Excel's limit."""
    cleaned = _ILLEGAL_XML_CHARACTERS.sub("", value)
    return cleaned[:_MAX_CELL_LENGTH]


def _safe_name(value: str) -> str:
    """Return a filesystem-safe fragment for an XLSX file name."""
    normalized = re.sub(r"[^0-9A-Za-zА-Яа-яЁё._-]+", "_", value).strip("._")
    return (normalized or "asset")[:120]
