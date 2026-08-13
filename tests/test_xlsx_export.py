import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from mpvm_jira.models import Asset, SoftwarePackage, Vulnerability
from mpvm_jira.jira.xlsx_export import XLSX_HEADERS, export_asset_xlsx


class XlsxExportTests(unittest.TestCase):
    def test_required_columns_and_missing_identifiers(self):
        asset = Asset(
            "asset-1",
            fqdn="srv.example.org",
            os_name="Debian",
            os_version="12",
            vulnerabilities=[
                Vulnerability(
                    cves=("CVE-2026-1234",),
                    packages=(
                        SoftwarePackage("openssl", "3.0.11-1"),
                        SoftwarePackage("libssl3", "3.0.11-1"),
                    ),
                    cvss_score=7.5,
                    severity="High",
                    cvss_vector="CVSS:3.1/AV:N/AC:L",
                    description="Описание, с запятой",
                    remediation="Установить обновление",
                )
            ],
        )
        with tempfile.TemporaryDirectory() as directory:
            path = export_asset_xlsx(asset, Path(directory))
            self.assertEqual(path.suffix, ".xlsx")
            workbook = load_workbook(path)
            sheet = workbook["Уязвимости"]
            rows = list(sheet.iter_rows(values_only=True))

        self.assertEqual(rows[0], XLSX_HEADERS)
        self.assertEqual(rows[1][1], "CVE-2026-1234")
        self.assertEqual(rows[1][2], "-")
        self.assertEqual(rows[1][3], "Debian 12")
        self.assertEqual(rows[1][4], "openssl; libssl3")
        self.assertEqual(rows[1][5], "3.0.11-1; 3.0.11-1")
        self.assertEqual(rows[1][7], "Установить обновление")
        self.assertEqual(
            rows[1][8:11],
            ("CVSS:3.1/AV:N/AC:L", "Высокий", 7.5),
        )

    def test_missing_values_are_written_as_dashes(self):
        asset = Asset("asset-1", vulnerabilities=[Vulnerability()])
        with tempfile.TemporaryDirectory() as directory:
            path = export_asset_xlsx(asset, Path(directory))
            workbook = load_workbook(path)
            sheet = workbook["Уязвимости"]
            row = tuple(cell.value for cell in sheet[2])
        self.assertEqual(row[3:6], ("-", "-", "-"))
        self.assertEqual(row[10], None)

    def test_workbook_is_formatted_for_excel(self):
        asset = Asset("asset-1", vulnerabilities=[Vulnerability()])
        with tempfile.TemporaryDirectory() as directory:
            path = export_asset_xlsx(asset, Path(directory))
            workbook = load_workbook(path)
            sheet = workbook["Уязвимости"]

        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet.auto_filter.ref, "A1:K2")
        self.assertTrue(sheet["A1"].font.bold)
        self.assertEqual(sheet["A1"].font.color.rgb, "00FFFFFF")
        self.assertTrue(sheet["G2"].alignment.wrap_text)
        self.assertEqual(sheet["K2"].number_format, "0.0")

    def test_invalid_xml_characters_do_not_break_export(self):
        vulnerability = Vulnerability(description="до\x00после")
        asset = Asset("asset-1", vulnerabilities=[vulnerability])
        with tempfile.TemporaryDirectory() as directory:
            path = export_asset_xlsx(asset, Path(directory))
            workbook = load_workbook(path)
            sheet = workbook["Уязвимости"]
            self.assertEqual(sheet["G2"].value, "допосле")

    def test_formula_like_text_is_stored_as_text(self):
        vulnerability = Vulnerability(description="=2+2")
        asset = Asset("asset-1", vulnerabilities=[vulnerability])
        with tempfile.TemporaryDirectory() as directory:
            path = export_asset_xlsx(asset, Path(directory))
            workbook = load_workbook(path, data_only=False)
            cell = workbook["Уязвимости"]["G2"]
            self.assertEqual(cell.value, "=2+2")
            self.assertEqual(cell.data_type, "s")


if __name__ == "__main__":
    unittest.main()
