import tempfile
import unittest
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from mpvm_jira.jira import IntegrationRunError, sync_latest_to_jira
from mpvm_jira.mpvm import save_snapshot


class FakeJiraClient:
    def __init__(self, *, fail_first_attachment=False):
        self.created = []
        self.created_assets = []
        self.created_criticalities = []
        self.attached = []
        self.validated = 0
        self.fail_first_attachment = fail_first_attachment

    def validate_connection(self):
        self.validated += 1
        return {"displayName": "Test"}

    def create_issue(self, asset, criticalities=None):
        self.created.append(asset.asset_id)
        self.created_assets.append(asset)
        self.created_criticalities.append(criticalities)
        return {"key": f"SEC-{len(self.created)}"}

    def attach_file(self, issue_key, path):
        self.attached.append((issue_key, path))
        if self.fail_first_attachment and len(self.attached) == 1:
            raise RuntimeError("temporary attachment error")
        return {"id": "42", "filename": path.name}


def payload():
    return {
        "schema_version": 2,
        "assets": [
            {
                "asset_id": "asset-1",
                "fqdn": "srv.example.org",
                "ip_addresses": ["192.0.2.10"],
                "importance": "Medium",
                "vulnerabilities": [
                    {
                        "vulnerability_id": "PT-1",
                        "cves": ["CVE-2026-1234"],
                        "bdus": ["BDU:2026-00001"],
                        "remediation": "Обновить пакет",
                    }
                ],
            }
        ],
    }


def criticality_payload():
    data = payload()
    vector = "CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:H/I:H/A:H"
    data["assets"][0]["vulnerabilities"] = [
        {
            "vulnerability_id": f"PT-{index}",
            "severity": severity,
            "cvss_vector": vulnerability_vector,
            "cvss_score": score,
        }
        for index, (severity, vulnerability_vector, score) in enumerate(
            (
                ("Critical", vector, 9.8),
                ("High", vector, 8.5),
                ("Medium", vector, 5.5),
                ("Low", vector, 2.1),
                ("Critical", "", 9.1),
            ),
            1,
        )
    ]
    return data


class ServiceTests(unittest.TestCase):
    def test_multiple_assets_create_one_issue_and_attachment_each(self):
        config = {
            "duplicate_policy": "skip_seen_snapshot",
            "state_file": ".state.json",
        }
        data = payload()
        second = deepcopy(data["assets"][0])
        second["asset_id"] = "asset-2"
        second["fqdn"] = "srv2.example.org"
        second["ip_addresses"] = ["192.0.2.11"]
        second["vulnerabilities"][0]["vulnerability_id"] = "PT-2"
        data["assets"].append(second)

        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            save_snapshot(base, data)
            client = FakeJiraClient()

            result = sync_latest_to_jira(client, base, config)

            self.assertEqual(client.created, ["asset-1", "asset-2"])
            self.assertEqual(len(result["created"]), 2)
            self.assertEqual(len(result["attached"]), 2)

    def test_same_snapshot_is_not_created_twice(self):
        config = {
            "duplicate_policy": "skip_seen_snapshot",
            "state_file": ".state.json",
        }
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            save_snapshot(
                base,
                payload(),
                datetime(2026, 8, 5, 12, 0, 0, tzinfo=timezone.utc),
            )
            client = FakeJiraClient()
            first = sync_latest_to_jira(client, base, config)
            second = sync_latest_to_jira(client, base, config)
            self.assertEqual(len(first["created"]), 1)
            self.assertEqual(len(first["attached"]), 1)
            self.assertEqual(second["created"], [])
            self.assertEqual(
                second["skipped"], ["srv.example.org (192.0.2.10)"]
            )
            self.assertEqual(len(client.created), 1)
            self.assertEqual(len(client.attached), 1)

    def test_attachment_retry_does_not_create_second_issue(self):
        config = {
            "duplicate_policy": "skip_seen_snapshot",
            "state_file": ".state.json",
        }
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            save_snapshot(
                base,
                payload(),
                datetime(2026, 8, 5, 12, 0, 0, tzinfo=timezone.utc),
            )
            client = FakeJiraClient(fail_first_attachment=True)
            with self.assertRaises(IntegrationRunError):
                sync_latest_to_jira(client, base, config)
            result = sync_latest_to_jira(client, base, config)
            self.assertEqual(len(client.created), 1)
            self.assertEqual(len(client.attached), 2)
            self.assertEqual(result["created"], [])
            self.assertEqual(result["attached"][0]["issue_key"], "SEC-1")

    def test_explicit_snapshot_is_used_instead_of_latest_file(self):
        config = {
            "duplicate_policy": "skip_seen_snapshot",
            "state_file": ".state.json",
        }
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            selected = save_snapshot(
                base,
                payload(),
                datetime(2026, 8, 5, 12, 0, 0, tzinfo=timezone.utc),
            )
            newer_payload = deepcopy(payload())
            newer_payload["assets"][0]["asset_id"] = "asset-2"
            save_snapshot(
                base,
                newer_payload,
                datetime(2026, 8, 5, 13, 0, 0, tzinfo=timezone.utc),
            )

            client = FakeJiraClient()
            sync_latest_to_jira(client, base, config, snapshot_path=selected)

            self.assertEqual(client.created, ["asset-1"])

    def test_filter_limits_issue_asset_and_xlsx_rows(self):
        config = {
            "duplicate_policy": "skip_seen_snapshot",
            "state_file": ".state.json",
        }
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            save_snapshot(base, criticality_payload())
            client = FakeJiraClient()

            result = sync_latest_to_jira(
                client,
                base,
                config,
                criticalities=["Critical", "High"],
            )

            self.assertEqual(len(client.created_assets), 1)
            self.assertEqual(
                len(client.created_assets[0].vulnerabilities),
                2,
            )
            self.assertEqual(
                client.created_criticalities,
                [("Critical", "High")],
            )
            xlsx_path = Path(result["attached"][0]["xlsx"])
            workbook = load_workbook(xlsx_path, read_only=True)
            rows = list(
                workbook["Уязвимости"].iter_rows(values_only=True)
            )
            self.assertEqual(len(rows), 3)
            self.assertEqual(
                [row[9] for row in rows[1:]], ["Критический", "Высокий"]
            )

    def test_no_matching_vulnerabilities_skips_jira_entirely(self):
        config = {
            "duplicate_policy": "skip_seen_snapshot",
            "state_file": ".state.json",
        }
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            save_snapshot(base, payload())
            client = FakeJiraClient()

            result = sync_latest_to_jira(
                client,
                base,
                config,
                criticalities=["Critical"],
            )

            self.assertEqual(client.validated, 0)
            self.assertEqual(client.created, [])
            self.assertEqual(
                result["filtered_out"],
                ["srv.example.org (192.0.2.10)"],
            )

    def test_different_filters_have_independent_duplicate_state(self):
        config = {
            "duplicate_policy": "skip_seen_snapshot",
            "state_file": ".state.json",
        }
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            save_snapshot(base, criticality_payload())
            client = FakeJiraClient()

            sync_latest_to_jira(
                client,
                base,
                config,
                criticalities=["Critical"],
            )
            sync_latest_to_jira(
                client,
                base,
                config,
                criticalities=["High"],
            )

            self.assertEqual(client.created, ["asset-1", "asset-1"])


if __name__ == "__main__":
    unittest.main()
